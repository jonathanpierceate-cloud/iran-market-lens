from __future__ import annotations

import asyncio
import csv
import io
import json
import math
import os
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import asynccontextmanager
from datetime import datetime, time as datetime_time, timedelta, timezone
from pathlib import Path
from threading import Lock, Thread
from typing import Any

from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from .analysis import analyze, analyze_rahavard, stale_status
from .config import (MAX_STALE_HOURS, RAHAVARD_API_BASE_URL,
                     RAHAVARD_ETF_FUNDS_URL, RAHAVARD_GOLD_PAGE_URL,
                     RAHAVARD_GOLD_18K_PAGE_URL,
                     RAHAVARD_LIGHT_BARS_URL, RAHAVARD_TEPIX_PAGE_URL,
                     RAHAVARD_USDT_PAGE_URL,
                     TABDEAL_USDT_PAGE_URL,
                     REFRESH_INTERVAL_SECONDS, ROOT)
from .db import db, utc_now
from .fund_scoring import FUND_TYPE_BY_CATEGORY_ID, FUND_TYPE_LABELS, resolve_fund_type
from .indicators import fund_chart_indicator_groups
from .scoring_profiles import normalize_horizon, horizon_label
from .scoring_rules import entry_recommendation
from .sources import (SourceError, fetch_fund_nav_history,
                      fetch_fund_price_history, fetch_rahavard_fund_profile,
                      rahavard_etf_fund_asset_ids,
                      rahavard_fund_history, refresh_all,
                      TABDEAL_USDT_CHART_PAGE_URL, tabdeal_usdt_history)

STATIC = ROOT / "app" / "static"
REFRESH_STATE: dict[str, Any] = {"running": False, "last_result": None}
REFRESH_WAKE: asyncio.Event | None = None
FUND_CATEGORY_CACHE: dict[str, tuple[float, set[str]]] = {}
FUND_CATEGORY_CACHE_TTL = 15 * 60
FUND_CATEGORY_MAP_CACHE: tuple[float, dict[str, str]] | None = None
FUND_CATEGORY_MAP_TTL = 2 * 60 * 60
FUND_CATEGORY_MAP_LOCK = Lock()
FUND_CATEGORY_MAP_BUILDING = False
DOLLAR_INDICATOR_CACHE: dict[str, Any] = {}
DOLLAR_INDICATOR_CACHE_TTL = 5 * 60


class WatchlistAddInput(BaseModel):
    symbol: str = Field(min_length=1, max_length=120)


class FundScoreBatchInput(BaseModel):
    fund_keys: list[str] = Field(default_factory=list, max_length=4)
    category_id: str = Field(default="", max_length=24)
    chart_bars: dict[str, list[dict[str, Any]]] | None = None
    chart_fetch_ms: dict[str, int] = Field(default_factory=dict)


class FundTechnicalBarsInput(BaseModel):
    bars: list[dict[str, Any]] = Field(min_length=2, max_length=500)
    chart_fetch_ms: int | None = Field(default=None, ge=0, le=120000)


class FundNavBatchInput(BaseModel):
    fund_keys: list[str] = Field(default_factory=list, max_length=8)


def _source_map() -> dict[str, dict[str, Any]]:
    return {row["name"]: row for row in db.source_status()}


def _market_history(symbol: str, limit: int = 1000) -> list[dict[str, Any]]:
    fund = db.get_fund(symbol)
    if fund and not fund.get("ambiguous") and fund.get("source") == "Rahavard365":
        return db.history(fund["fund_key"], limit, source="Rahavard365")
    preferred = {"GOLD": "Rahavard365", "GOLD_18K": "Rahavard365", "USD_IR_FREE": "Rahavard365", "TEPIX": "Rahavard365"}.get(symbol)
    if preferred:
        bars = db.history(symbol, limit, source=preferred)
        if symbol == "USD_IR_FREE":
            # Do not mix candles from the former Rahavard asset with the requested USDT asset.
            bars = [bar for bar in bars if bar.get("source_url") == RAHAVARD_USDT_PAGE_URL]
            if not bars:
                return []
        if bars:
            if symbol == "USD_IR_FREE":
                iran_tz = timezone(timedelta(hours=3, minutes=30))
                daily: dict[str, dict[str, Any]] = {}
                for bar in bars:
                    try:
                        stamp = datetime.fromisoformat(str(bar["timestamp"]).replace("Z", "+00:00"))
                        if stamp.tzinfo is None:
                            stamp = stamp.replace(tzinfo=timezone.utc)
                        day = stamp.astimezone(iran_tz).date().isoformat()
                    except (TypeError, ValueError):
                        continue
                    price = bar.get("close")
                    if price is None:
                        continue
                    current = daily.get(day)
                    if current is None:
                        daily[day] = {**bar, "open": price, "high": price, "low": price, "volume": None}
                    else:
                        current["high"] = max(current["high"], price)
                        current["low"] = min(current["low"], price)
                        current["close"] = price
                        current["timestamp"] = bar["timestamp"]
                        current["fetched_at"] = bar.get("fetched_at")
                bars = [daily[day] for day in sorted(daily)]
                return bars[-max(1, limit):]
            return bars
    return db.history(symbol, limit)


def _tabdeal_dollar_chart_history() -> tuple[list[dict[str, Any]], str]:
    """Load Tabdeal's daily USDT/IRT candles for dashboard indicator calculations."""
    cache_key = "Tabdeal:USDT_IRT:1D"
    cached = DOLLAR_INDICATOR_CACHE.get(cache_key)
    cache_age = time.monotonic() - cached["cached_at"] if cached else None
    if cached and cache_age is not None and cache_age < DOLLAR_INDICATOR_CACHE_TTL:
        bars = cached["bars"]
        timestamp = bars[-1].get("timestamp") if bars else None
        return bars, ("stale" if timestamp and stale_status(timestamp, max_hours=MAX_STALE_HOURS)
                      else "available")

    try:
        bars = tabdeal_usdt_history()
        DOLLAR_INDICATOR_CACHE[cache_key] = {
            "bars": bars, "cached_at": time.monotonic(),
        }
        timestamp = bars[-1].get("timestamp") if bars else None
        return bars, ("stale" if timestamp and stale_status(timestamp, max_hours=MAX_STALE_HOURS)
                      else "available")
    except Exception:
        if cached and cached.get("bars"):
            bars = cached["bars"]
            timestamp = bars[-1].get("timestamp") if bars else None
            return bars, ("stale" if timestamp and stale_status(timestamp, max_hours=MAX_STALE_HOURS)
                          else "cached_after_source_error")
        return [], "failed"


def _analyze_chart_candles(symbol: str, bars: list[dict[str, Any]], *,
                           horizon: str, indicator_source: str,
                           source_url: str | None = None) -> dict[str, Any]:
    """Use the same locally calculated, four-family analysis pipeline as funds."""
    timestamp = bars[-1].get("timestamp") if bars else None
    payload = fund_chart_indicator_groups(bars, source=indicator_source)
    context = {"fund_history": bars, "price_timestamp": timestamp}
    current_price = bars[-1].get("close") if bars else None
    result = analyze_rahavard(
        symbol, payload, timestamp, current_price, context,
        horizon=horizon, indicator_source=indicator_source)
    result["indicator_source_url"] = (
        source_url or (bars[-1].get("source_url") if bars else None))
    result["indicator_data_timestamp"] = timestamp
    result["price_is_cached"] = _timestamp_is_cached(timestamp)
    return result


def _refresh_interval() -> int:
    try:
        return max(60, min(86400, int(db.get_settings().get("refresh_interval_seconds", REFRESH_INTERVAL_SECONDS))))
    except (TypeError, ValueError):
        return REFRESH_INTERVAL_SECONDS


def _market_quote(symbol: str, horizon: str = "long") -> dict[str, Any]:
    is_usd_quote = symbol == "USD_IR_FREE"
    bars = _market_history(symbol, 10000)
    quote_bars = db.history(symbol, 10, source="Tabdeal") if is_usd_quote else bars
    quote_bar = quote_bars[-1] if quote_bars else (bars[-1] if bars and not is_usd_quote else None)
    history_latest = bars[-1] if bars else None
    with db.connect() as conn:
        asset = conn.execute("SELECT * FROM assets WHERE symbol=?", (symbol,)).fetchone()
    asset_defaults = {
        "GOLD": {"name": "طلای جهانی", "source": "Rahavard365", "source_url": RAHAVARD_GOLD_PAGE_URL},
        "GOLD_18K": {"name": "طلای ۱۸ عیار", "source": "Rahavard365", "source_url": RAHAVARD_GOLD_18K_PAGE_URL},
        "USD_IR_FREE": {"name": "دلار", "source": "Tabdeal", "source_url": TABDEAL_USDT_PAGE_URL},
        "TEPIX": {"name": "شاخص کل بورس", "source": "Rahavard365", "source_url": RAHAVARD_TEPIX_PAGE_URL},
    }
    data = dict(asset) if asset else asset_defaults.get(symbol)
    if data is None:
        raise HTTPException(status_code=404, detail="دارایی پیدا نشد")
    display_name = asset_defaults.get(symbol, {}).get("name") or data["name"]
    source_url = TABDEAL_USDT_PAGE_URL if is_usd_quote else data.get("source_url")
    source_key = {"GOLD": "rahavard_gold", "GOLD_18K": "rahavard_gold_18k", "DXY": "yahoo_dxy",
                  "USD_IR_FREE": "tabdeal_usdt", "TEPIX": "rahavard_tepix"}.get(symbol)
    quote_source = _source_map().get(source_key) if source_key else None
    analysis_source_key = "tabdeal_usdt" if is_usd_quote else source_key
    analysis_source = _source_map().get(analysis_source_key) if analysis_source_key else None
    analysis_bars = bars
    chart_analysis_status = None
    if is_usd_quote:
        analysis_bars, chart_analysis_status = _tabdeal_dollar_chart_history()

    history_returns = {}
    for label, periods in [("daily", 1), ("weekly", 5), ("monthly", 21), ("three_month", 63), ("six_month", 126), ("one_year", 252)]:
        if len(bars) > periods and bars[-periods-1].get("close") and history_latest:
            history_returns[label] = round((history_latest["close"] / bars[-periods-1]["close"] - 1) * 100, 3)
        else:
            history_returns[label] = None

    price = quote_bar.get("close") if quote_bar else None
    returns = {}
    for label, periods in [("daily", 1), ("weekly", 5), ("monthly", 21), ("three_month", 63), ("six_month", 126), ("one_year", 252)]:
        baseline = bars[-periods-1].get("close") if len(bars) > periods else None
        if is_usd_quote and baseline:
            # Tabdeal quotes toman; Rahavard's dollar candles are stored in rial.
            baseline /= 10
        returns[label] = round((price / baseline - 1) * 100, 3) if price is not None and baseline else None

    observed_timestamp = ((quote_source or {}).get("last_data_timestamp")
                         or (quote_bar or {}).get("timestamp"))
    analysis_timestamp = ((analysis_bars[-1].get("timestamp") if analysis_bars else None)
                          if is_usd_quote else
                          (analysis_source or {}).get("last_data_timestamp")
                          or (history_latest or {}).get("timestamp"))
    price_stale = stale_status(observed_timestamp, max_hours=MAX_STALE_HOURS) if observed_timestamp else True
    analysis_stale = stale_status(analysis_timestamp, max_hours=MAX_STALE_HOURS) if analysis_timestamp else True

    def status_for(feed: dict[str, Any] | None, stale: bool) -> str:
        if feed and feed.get("status") == "stale":
            return "stale"
        if feed and feed.get("status") == "failed":
            return "stale" if stale else "cached_after_source_error"
        return "stale" if stale else "available"

    status = status_for(quote_source, price_stale) if price is not None else "unavailable"
    output = {
        "symbol": symbol, "name": display_name, "price": price,
        "open": quote_bar.get("open") if quote_bar else None,
        "high": quote_bar.get("high") if quote_bar else None,
        "low": quote_bar.get("low") if quote_bar else None,
        "volume": quote_bar.get("volume") if quote_bar else None,
        "history": bars, "returns": returns,
        "source": "Tabdeal" if is_usd_quote and price is not None else (quote_bar.get("source") if quote_bar else data.get("source")),
        "source_url": source_url, "data_timestamp": observed_timestamp,
        "fetched_at": quote_bar.get("fetched_at") if quote_bar else None,
        "stale": bool(price_stale), "status": status,
    }
    output["price_is_cached"] = _timestamp_is_cached(observed_timestamp)
    chart_indicator_source = (
        "TabdealChart" if is_usd_quote else
        "Rahavard365Chart" if output.get("source") == "Rahavard365" else None)
    if chart_indicator_source:
        output["analysis"] = _analyze_chart_candles(
            symbol, analysis_bars, horizon=horizon,
            indicator_source=chart_indicator_source,
            source_url=(TABDEAL_USDT_CHART_PAGE_URL if is_usd_quote else None))
    else:
        output["analysis"] = analyze(symbol, analysis_bars, horizon=horizon)
    output["analysis"]["source"] = "Tabdeal" if is_usd_quote else output.get("source")
    output["analysis"]["data_timestamp"] = analysis_timestamp
    output["analysis"]["status"] = (chart_analysis_status if is_usd_quote
                                      else status_for(analysis_source, analysis_stale))
    if is_usd_quote or output.get("source") == "Rahavard365":
        output["analysis"]["indicator_source"] = chart_indicator_source
        output["analysis"]["indicator_source_url"] = (
            TABDEAL_USDT_CHART_PAGE_URL if is_usd_quote else
            (analysis_bars[-1].get("source_url") if analysis_bars else data.get("source_url"))
        )
        output["analysis"]["indicator_data_timestamp"] = analysis_timestamp
    if not is_usd_quote:
        output["analysis"]["daily_return_pct"] = history_returns.get("daily")
    output["analysis"]["price_is_cached"] = _timestamp_is_cached(analysis_timestamp)
    return output


def evaluate_alerts() -> list[dict[str, Any]]:
    """Evaluate active local alerts only against fresh source-backed observations."""
    triggered = []
    for alert in db.alerts():
        if not alert.get("enabled") or alert.get("triggered_at"):
            continue
        fund = db.get_fund(alert["symbol"])
        if fund and fund.get("ambiguous"):
            continue
        if fund and fund.get("source") == "Rahavard365" and alert["condition"] in {
                "rsi_overbought", "rsi_oversold", "volume_spike"}:
            raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
            asset_id = str(raw.get("rahavard_asset_id") or fund.get("registration_no") or fund.get("symbol") or "").strip()
            try:
                bars = rahavard_fund_history(asset_id)
            except SourceError:
                continue
            if alert["condition"] == "volume_spike":
                # Volume alerts remain disabled until their configured rule is defined.
                continue
            if stale_status(bars[-1].get("timestamp"), max_hours=MAX_STALE_HOURS):
                continue
            indicator_payload = fund_chart_indicator_groups(bars)
            result = analyze_rahavard(
                fund["symbol"], indicator_payload, bars[-1].get("timestamp"),
                bars[-1].get("close"), indicator_source="Rahavard365Chart")
            rsi_item = next((item for item in result["indicators"].get("oscillators", [])
                             if str(item.get("short_name_en") or "").startswith("RSI(14)")), None)
            rsi_value = None
            if rsi_item:
                for part in rsi_item.get("value") or []:
                    if str(part.get("name") or "").casefold() in {"rsi", "value"}:
                        try:
                            rsi_value = float(part.get("value"))
                        except (TypeError, ValueError):
                            pass
                        break
            hit = ((alert["condition"] == "rsi_overbought" and rsi_value is not None and rsi_value >= 70)
                   or (alert["condition"] == "rsi_oversold" and rsi_value is not None and rsi_value <= 30))
            if hit:
                reason = f"RSI(14) ره‌آورد به {rsi_value:.1f} رسید"
                db.mark_alert_triggered(alert["id"], reason)
                triggered.append({"id": alert["id"], "symbol": fund["symbol"], "message": reason})
            continue
        history_symbol = fund["fund_key"] if fund else alert["symbol"]
        analysis_symbol = fund["symbol"] if fund else alert["symbol"]
        bars = _market_history(history_symbol, 5000)
        if not bars or stale_status(bars[-1].get("timestamp"), max_hours=MAX_STALE_HOURS):
            continue
        current = bars[-1].get("close")
        condition = alert["condition"]
        hit = False
        reason = ""
        if condition == "price_above" and current is not None and alert.get("threshold") is not None:
            hit = current >= alert["threshold"]; reason = f"قیمت به {current} رسید و از آستانه گذشت"
        elif condition == "price_below" and current is not None and alert.get("threshold") is not None:
            hit = current <= alert["threshold"]; reason = f"قیمت به {current} رسید و از آستانه پایین‌تر آمد"
        elif condition in {"rsi_overbought", "rsi_oversold", "volume_spike"}:
            result = analyze(analysis_symbol, bars)
            indicators = result["indicators"]["indicators"]
            rsi = indicators.get("RSI(14)", {}).get("value")
            spike = result["indicators"].get("volume_analysis", {}).get("volume_spike")
            if condition == "rsi_overbought": hit = rsi is not None and rsi >= 70; reason = f"RSI(14) به {rsi:.1f} رسید" if hit else ""
            elif condition == "rsi_oversold": hit = rsi is not None and rsi <= 30; reason = f"RSI(14) به {rsi:.1f} رسید" if hit else ""
            elif condition == "volume_spike": hit = bool(spike); reason = "جهش حجم نسبت به میانگین ۲۰ دوره ثبت شد" if hit else ""
        if hit:
            db.mark_alert_triggered(alert["id"], reason)
            triggered.append({"id": alert["id"], "symbol": analysis_symbol, "message": reason})
    return triggered


async def _periodic_refresh():
    first_run = True
    while True:
        if not first_run:
            if REFRESH_WAKE is None:
                await asyncio.sleep(_refresh_interval())
            else:
                try:
                    await asyncio.wait_for(REFRESH_WAKE.wait(), timeout=_refresh_interval())
                except asyncio.TimeoutError:
                    pass
                else:
                    REFRESH_WAKE.clear()
                    continue
        first_run = False
        try:
            if not REFRESH_STATE["running"]:
                REFRESH_STATE["running"] = True
                REFRESH_STATE["last_result"] = await refresh_all()
                evaluate_alerts()
        except Exception as exc:
            REFRESH_STATE["last_result"] = {"error": str(exc)}
        finally:
            REFRESH_STATE["running"] = False
@asynccontextmanager
async def lifespan(_: FastAPI):
    global REFRESH_WAKE
    db.init()
    _schedule_fund_category_map_refresh()
    REFRESH_WAKE = asyncio.Event()
    for key, label, url in [
        ("rahavard_gold", "طلای جهانی / ره‌آورد۳۶۵", RAHAVARD_GOLD_PAGE_URL),
        ("rahavard_gold_18k", "طلای ۱۸ عیار / ره‌آورد۳۶۵", RAHAVARD_GOLD_18K_PAGE_URL),
        ("rahavard_usdt", "تاریخچهٔ دلار / ره‌آورد۳۶۵", RAHAVARD_USDT_PAGE_URL),
        ("tabdeal_usdt", "قیمت دلار / تتر · Tabdeal", TABDEAL_USDT_PAGE_URL),
        ("rahavard_tepix", "شاخص کل بورس / ره‌آورد۳۶۵", RAHAVARD_TEPIX_PAGE_URL),
        ("yahoo_dxy", "شاخص دلار آمریکا", "https://finance.yahoo.com/quote/DX-Y.NYB/"),
        ("rahavard_funds", "صندوق‌های قابل معامله / ره‌آورد۳۶۵", RAHAVARD_ETF_FUNDS_URL),
        ("rahavard_fund_details", "نمایه و NAV صندوق / ره‌آورد۳۶۵", RAHAVARD_API_BASE_URL + "/asset/{asset_id}"),
        ("rahavard_fund_charts", "نمودار و محاسبهٔ تکنیکال صندوق / ره‌آورد۳۶۵", RAHAVARD_LIGHT_BARS_URL + "?symbol=exchange.asset:{asset_id}:real_close"),
    ]:
        existing_source = _source_map().get(key)
        if existing_source is None:
            db.save_source(key, label, "not_tested", url, None, "هنوز بررسی نشده است", None)
        elif key == "tabdeal_usdt" and existing_source.get("status") == "disabled":
            db.save_source(key, label, "not_tested", url, None, None, None)
        elif existing_source.get("label") != label:
            db.update_source_label(key, label)
    legacy_sources = _source_map()
    for key, label in [("tsetmc", "منبع قبلی قیمت صندوق‌ها · غیرفعال"),
                       ("tsetmc_fund_technicals", "تحلیل تکنیکال صندوق‌ها / TSETMC · غیرفعال"),
                       ("fipiran", "منبع قبلی مشخصات صندوق‌ها · غیرفعال"),
                       ("rahavard_indicators", "اندیکاتورهای قبلی صندوق‌ها / ره‌آورد · غیرفعال")]:
        previous = legacy_sources.get(key)
        if previous and previous.get("status") != "disabled":
            db.save_source(key, label, "disabled", previous.get("source_url"),
                           previous.get("last_data_timestamp"),
                            "تحلیل صندوق‌ها اکنون از کندل‌های نمودار ره‌آورد محاسبه می‌شود.", None)
        elif previous and previous.get("label") != label:
            db.update_source_label(key, label)
    for key, label in [("yahoo_gold", "Yahoo Finance · تاریخچه پیشین"),
                       ("tgju", "TGJU · تاریخچه پیشین")]:
        previous = legacy_sources.get(key)
        if previous:
            db.save_source(key, label, "not_tested", previous.get("source_url"),
                           previous.get("last_data_timestamp"),
                           "این منبع برای داده‌های جدید استفاده نمی‌شود؛ تاریخچه قبلی نگهداری شده است.", None)
    # Vercel may freeze a function immediately after a response.  A forever
    # running task is therefore not a reliable scheduler there; refreshes are
    # triggered explicitly through /api/refresh (and can be called by a cron).
    if os.getenv("VERCEL"):
        # A new Vercel instance has an empty /tmp directory.  Bootstrap it
        # once before serving the first request so a cold instance still has
        # live market prices, fund listings, and NAV instead of an empty
        # dashboard.  Warm instances can be refreshed through /api/refresh.
        if not db.funds(source="Rahavard365") or not db.history("GOLD_18K", 1, source="Rahavard365"):
            REFRESH_STATE["running"] = True
            try:
                REFRESH_STATE["last_result"] = await refresh_all()
                evaluate_alerts()
            except Exception as exc:
                REFRESH_STATE["last_result"] = {"error": str(exc)}
            finally:
                REFRESH_STATE["running"] = False
        yield
        return

    task = asyncio.create_task(_periodic_refresh())
    yield
    task.cancel()
    try:
        await task
    except asyncio.CancelledError:
        pass


app = FastAPI(title="AliVest | تحلیل بازار سرمایه ایران", version="1.0.0", lifespan=lifespan,
              description="داشبورد محلی با ثبت منبع، زمان داده و وضعیت تازگی.")
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/", response_class=HTMLResponse)
def home():
    return (STATIC / "index.html").read_text(encoding="utf-8")


@app.get("/api/health")
def health():
    try:
        with db.connect() as conn:
            conn.execute("SELECT 1")
            counts = {table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                      for table in ["prices", "funds", "alerts"]}
        counts["funds"] = len(db.funds(source="Rahavard365"))
        return {"status": "ok", "database": "connected", "counts": counts,
                "scheduler": "disabled_on_vercel" if os.getenv("VERCEL") else "running",
                "checked_at": utc_now()}
    except Exception as exc:
        return JSONResponse(status_code=503, content={"status": "error", "database": str(exc)})


def _gold18k_valuation(gold_quote: dict[str, Any], dollar_quote: dict[str, Any], rahavard_quote: dict[str, Any]) -> dict[str, Any]:
    """Estimate 18K gold per gram from XAU/USD and Tabdeal USDT price in toman."""
    def number(value: Any) -> float | None:
        try:
            result = float(value)
        except (TypeError, ValueError):
            return None
        return result if math.isfinite(result) and result > 0 else None

    gold_usd_per_ounce = number(gold_quote.get("price"))
    usdt_toman = number(dollar_quote.get("price"))
    rahavard_irr_per_gram = number(rahavard_quote.get("price"))
    valuation: dict[str, Any] = {
        "theoretical_price": None,
        "rahavard_price": rahavard_irr_per_gram,
        "bubble": None,
        "bubble_percent": None,
        "gold_usd_per_ounce": gold_usd_per_ounce,
        "usdt_toman": usdt_toman,
        "purity": 0.75,
        "ounce_grams": 31.1034768,
        "method": "XAU/USD * USDT(Toman) * 10 * 75% / 31.1034768",
    }
    if gold_usd_per_ounce is None or usdt_toman is None or rahavard_irr_per_gram is None:
        return valuation

    # Tabdeal quotes USDT in toman; the Rahavard 18K gold feed is in rial.
    theoretical = gold_usd_per_ounce * usdt_toman * 10 * 0.75 / 31.1034768
    bubble = rahavard_irr_per_gram - theoretical
    valuation.update({
        "theoretical_price": round(theoretical),
        "bubble": round(bubble),
        "bubble_percent": round((bubble / theoretical) * 100, 2) if theoretical else None,
    })
    return valuation


@app.get("/api/market/summary")
def market_summary(horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    funds = _exclude_fixed_income_funds(db.funds(sort="data_timestamp", source="Rahavard365"))
    gold_quote = _market_quote("GOLD", horizon)
    gold18k_quote = _market_quote("GOLD_18K", horizon)
    dollar_quote = _market_quote("USD_IR_FREE", horizon)
    tepix_quote = _market_quote("TEPIX", horizon)
    gold_quote.pop("history", None)
    gold18k_quote.pop("history", None)
    dollar_quote.pop("history", None)
    tepix_quote.pop("history", None)
    gold18k_quote["valuation"] = _gold18k_valuation(gold_quote, dollar_quote, gold18k_quote)
    return {"as_of": utc_now(), "investment_horizon": horizon,
            "investment_horizon_label": horizon_label(horizon),
            "gold": gold_quote, "gold18k": gold18k_quote,
            "dollar": dollar_quote,
            "tepix": tepix_quote,
            "funds_total": len(funds),
            "top_funds": sorted(funds, key=lambda f: f.get("value") or 0, reverse=True)[:5],
            "sources": db.source_status(), "refresh": REFRESH_STATE}


@app.get("/api/gold")
def gold(horizon: str = Query(default="long")):
    return _market_quote("GOLD", normalize_horizon(horizon))


@app.get("/api/gold-18k")
def gold_18k(horizon: str = Query(default="long")):
    return _market_quote("GOLD_18K", normalize_horizon(horizon))


@app.get("/api/gold/history")
def gold_history(limit: int = Query(default=1000, ge=1, le=5000)):
    bars = _market_history("GOLD", limit)
    source = _source_map().get("rahavard_gold")
    observed = (source or {}).get("last_data_timestamp") or (bars[-1]["timestamp"] if bars else None)
    return {"symbol": "GOLD", "history": bars,
            "source": bars[-1]["source"] if bars else None,
            "data_timestamp": observed,
            "stale": stale_status(observed, max_hours=MAX_STALE_HOURS) if observed else None}


@app.get("/api/dollar")
def dollar(horizon: str = Query(default="long")):
    return _market_quote("USD_IR_FREE", normalize_horizon(horizon))


@app.get("/api/dollar/indicators")
async def dollar_indicators(horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    bars, status = await asyncio.to_thread(_tabdeal_dollar_chart_history)
    if not bars:
        return {
            "analysis": None, "source": "TabdealChart",
            "source_url": TABDEAL_USDT_CHART_PAGE_URL,
            "status": "failed", "data_timestamp": None,
            "error": "کندل روزانهٔ دلار در نمودار Tabdeal برای محاسبهٔ اندیکاتورها در دسترس نیست.",
        }
    data_timestamp = bars[-1].get("timestamp")
    analysis = _analyze_chart_candles(
        "USD_IR_FREE", bars, horizon=horizon,
        indicator_source="TabdealChart", source_url=TABDEAL_USDT_CHART_PAGE_URL)
    analysis["source"] = "Tabdeal"
    analysis["data_timestamp"] = data_timestamp
    analysis["status"] = status
    analysis["indicator_source"] = "TabdealChart"
    analysis["indicator_source_url"] = TABDEAL_USDT_CHART_PAGE_URL
    analysis["indicator_data_timestamp"] = data_timestamp
    return {
        "analysis": analysis, "source": "TabdealChart",
        "source_url": TABDEAL_USDT_CHART_PAGE_URL,
        "status": status, "data_timestamp": data_timestamp,
    }


@app.get("/api/tepix")
def tepix(horizon: str = Query(default="long")):
    return _market_quote("TEPIX", normalize_horizon(horizon))


@app.get("/api/dollar/history")
def dollar_history(limit: int = Query(default=1000, ge=1, le=5000)):
    bars = _market_history("USD_IR_FREE", limit)
    source = _source_map().get("rahavard_usdt")
    observed = (source or {}).get("last_data_timestamp") or (bars[-1]["timestamp"] if bars else None)
    return {"symbol": "USD_IR_FREE", "history": bars,
            "source": bars[-1]["source"] if bars else None,
            "data_timestamp": observed,
            "stale": stale_status(observed, max_hours=MAX_STALE_HOURS) if observed else None}


@app.get("/api/market/sources")
def data_sources():
    return {"sources": db.source_status(), "checked_at": utc_now()}


@app.api_route("/api/refresh", methods=["GET", "POST"])
async def manual_refresh():
    if REFRESH_STATE["running"]:
        return {"status": "already_running", "last_result": REFRESH_STATE["last_result"]}
    REFRESH_STATE["running"] = True
    try:
        result = await refresh_all()
        REFRESH_STATE["last_result"] = result
        evaluate_alerts()
        return result
    finally:
        REFRESH_STATE["running"] = False


def _fund_observation_rank(fund: dict[str, Any]) -> tuple[float, int]:
    stamp = fund.get("data_timestamp")
    try:
        observed = datetime.fromisoformat(str(stamp).replace("Z", "+00:00"))
        if observed.tzinfo is None:
            observed = observed.replace(tzinfo=timezone.utc)
        timestamp = observed.timestamp()
    except (TypeError, ValueError, OverflowError):
        timestamp = float("-inf")
    return timestamp, int(fund.get("fund_key") == db.fund_key(fund))


def _fund_updated_at(fund: dict[str, Any]) -> str | None:
    raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
    return raw.get("trade_date_time") or raw.get("data_updated_at") or fund.get("data_timestamp")


def _timestamp_is_cached(timestamp: str | None) -> bool:
    """Treat missing or non-current Tehran trading-day data as cached."""
    if not timestamp:
        return True
    iran_tz = timezone(timedelta(hours=3, minutes=30))
    try:
        observed_day = datetime.fromisoformat(str(timestamp).replace("Z", "+00:00"))
        if observed_day.tzinfo is None:
            observed_day = observed_day.replace(tzinfo=iran_tz)
        return observed_day.astimezone(iran_tz).date() != datetime.now(iran_tz).date()
    except (TypeError, ValueError, OverflowError):
        return True


def _fund_category_id(fund: dict[str, Any], requested: str | None = None) -> str | None:
    raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
    for value in (requested, fund.get("category_id"), raw.get("groupId"),
                  raw.get("group_id"), raw.get("category_id"), raw.get("categoryId")):
        candidate = str(value or "").strip()
        if candidate in FUND_TYPE_BY_CATEGORY_ID:
            return candidate
    return None


def _fund_peer_group_key(fund: dict[str, Any]) -> str | None:
    category_id = _fund_category_id(fund)
    kind = resolve_fund_type({
        "category_id": category_id, "category": fund.get("category"),
        "fund_type": fund.get("fund_type"), "name": fund.get("name"),
    })
    return category_id or (f"type:{kind}" if kind else None)


def _record_fund_score_history(result: dict[str, Any] | None, fund_key: str,
                               data_timestamp: str | None,
                               price_history: list[dict[str, Any]] | None = None) -> None:
    if not isinstance(result, dict) or result.get("technical_score") is None:
        return
    score_date = str(data_timestamp or result.get("score_timestamp") or "")
    method = str(result.get("score_method") or "")
    try:
        db.save_fund_score_snapshot(
            fund_key, score_date, float(result["technical_score"]),
            result.get("data_coverage_pct"), method)
        bars = price_history if price_history is not None else db.history(
            fund_key, 60, source="Rahavard365")
        sessions = [str(row.get("timestamp") or "")[:10] for row in bars]
        result.update(db.fund_score_changes(fund_key, score_date, method, sessions))
    except Exception:
        # Some serverless deployments do not provide durable writable storage.
        result["score_change_5d"] = None
        result["score_change_20d"] = None


def _fund_model_context(fund: dict[str, Any], *, category_id: str | None = None,
                        category: Any = None, price_history: list[dict[str, Any]] | None = None,
                        nav_timestamp: str | None = None,
                        peer_items: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
    category_obj = category if isinstance(category, dict) else {}
    category_name = category_obj.get("name") if category_obj else category
    category_id = _fund_category_id(fund, str(category_id or category_obj.get("id") or "")) or ""
    context = {
        "category_id": category_id,
        "category": category_name or fund.get("category"),
        "fund_type": fund.get("fund_type"),
        "name": fund.get("name"),
        "symbol": fund.get("symbol"),
        "nav": _watchlist_number(fund.get("nav")),
        "nav_premium_pct": _watchlist_number(fund.get("nav_premium_pct")),
        "nav_timestamp": nav_timestamp,
        "price_timestamp": fund.get("data_timestamp") or raw.get("trade_date_time"),
        "daily_return_pct": _watchlist_number(fund.get("daily_return")),
        "fund_history": price_history,
    }
    kind = resolve_fund_type(context)
    context["resolved_fund_type"] = kind
    if kind:
        context["fund_type"] = kind
    if context["nav_timestamp"] is None and kind is not None:
        try:
            nav_rows = db.fund_nav_history(str(fund.get("fund_key") or ""))
            if nav_rows:
                latest_nav = nav_rows[-1]
                context["nav_timestamp"] = latest_nav.get("timestamp")
                context["nav"] = context["nav"] or _watchlist_number(latest_nav.get("statistical_nav"))
        except Exception:
            pass
    if context["fund_history"] is None:
        try:
            context["fund_history"] = _market_history(str(fund.get("fund_key") or fund.get("symbol") or ""), 150)
        except Exception:
            context["fund_history"] = []

    if peer_items:
        key = str(fund.get("fund_key") or "")
        liquidity: list[tuple[str, float]] = []
        for item in peer_items:
            traded_value = _watchlist_number(item.get("value"))
            if traded_value is not None and traded_value >= 0:
                liquidity.append((str(item.get("fund_key") or ""), traded_value))
        current_liquidity = next((value for peer_key, value in liquidity if peer_key == key),
                                 _watchlist_number(fund.get("value")))
        if current_liquidity is not None and len(liquidity) >= 5:
            below = sum(value < current_liquidity for _, value in liquidity)
            equal = sum(value == current_liquidity for _, value in liquidity)
            context["liquidity_percentile"] = round((below + equal / 2) / len(liquidity) * 100, 1)
    return context


async def _analyze_fund_chart(fund: dict[str, Any], *,
                              horizon: str = "long",
                              profile: dict[str, Any] | None = None,
                              model_context: dict[str, Any] | None = None,
                              preloaded_bars: list[dict[str, Any]] | None = None
                              ) -> tuple[dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    profile = profile if isinstance(profile, dict) else {}
    asset = profile.get("asset") if isinstance(profile.get("asset"), dict) else {}
    raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
    bars = (await fetch_fund_price_history(fund["fund_key"])
            if preloaded_bars is None else _normalize_client_chart_bars(preloaded_bars))
    if not bars:
        raise SourceError("کندل‌های نمودار ره‌آورد برای محاسبهٔ تحلیل صندوق موجود نیست")
    payload = fund_chart_indicator_groups(bars)
    technical_price = float(bars[-1]["close"])
    context = dict(model_context or _fund_model_context(fund, price_history=bars))
    context["fund_history"] = bars
    context["price_timestamp"] = bars[-1].get("timestamp")
    trade_symbol = (asset.get("trade_symbol") or raw.get("trade_symbol")
                    or raw.get("name") or fund.get("symbol"))
    result = analyze_rahavard(
        trade_symbol, payload,
        bars[-1].get("timestamp"), technical_price, context,
        horizon=horizon, indicator_source="Rahavard365Chart")
    result["indicator_source_url"] = bars[-1].get("source_url")
    result["indicator_data_timestamp"] = bars[-1].get("timestamp")
    result["price_is_cached"] = _timestamp_is_cached(bars[-1].get("timestamp"))
    return result, bars, payload


def _normalize_client_chart_bars(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_date: dict[str, dict[str, Any]] = {}
    for row in rows:
        if not isinstance(row, dict) or str(row.get("source") or "") != "Rahavard365":
            continue
        timestamp = str(row.get("timestamp") or "")[:10]
        source_url = str(row.get("source_url") or "")
        if (len(timestamp) != 10 or timestamp[4:5] != "-"
                or not source_url.startswith(RAHAVARD_LIGHT_BARS_URL)):
            continue
        try:
            close = float(row.get("close"))
            open_price = float(row.get("open", close))
            high = float(row.get("high", close))
            low = float(row.get("low", close))
            volume = float(row.get("volume") or 0)
        except (TypeError, ValueError, OverflowError):
            continue
        if not all(math.isfinite(value) for value in (close, open_price, high, low, volume)) or close <= 0:
            continue
        by_date[timestamp] = {
            "timestamp": timestamp, "fetched_at": row.get("fetched_at"),
            "open": open_price, "high": high, "low": low, "close": close,
            "volume": max(0.0, volume), "source": "Rahavard365", "source_url": source_url,
            "asset_id": str(row.get("asset_id") or ""),
        }
    bars = [by_date[key] for key in sorted(by_date)][-500:]
    if len(bars) < 2:
        raise SourceError("کندل معتبر نمودار ره‌آورد به محاسبه‌گر نرسید")
    return bars


def _fund_category_asset_ids(category_id: str) -> set[str]:
    now = time.monotonic()
    cached = FUND_CATEGORY_CACHE.get(category_id)
    if cached and now - cached[0] < FUND_CATEGORY_CACHE_TTL:
        return cached[1]
    asset_ids = rahavard_etf_fund_asset_ids(category_id)
    FUND_CATEGORY_CACHE[category_id] = (now, asset_ids)
    return asset_ids


def _fund_category_map_is_fresh(cache: tuple[float, dict[str, str]] | None,
                                now: float | None = None) -> bool:
    if not cache:
        return False
    now = time.monotonic() if now is None else now
    ttl = FUND_CATEGORY_MAP_TTL if cache[1] else 60
    return now - cache[0] < ttl


def _fund_category_map() -> dict[str, str]:
    """Build a short-lived asset-to-type map from Rahavard category lists."""
    global FUND_CATEGORY_MAP_CACHE
    now = time.monotonic()
    with FUND_CATEGORY_MAP_LOCK:
        now = time.monotonic()
        if _fund_category_map_is_fresh(FUND_CATEGORY_MAP_CACHE, now):
            return FUND_CATEGORY_MAP_CACHE[1]
        category_ids = tuple(FUND_TYPE_BY_CATEGORY_ID)
        mapping: dict[str, str] = {}
        with ThreadPoolExecutor(max_workers=8) as pool:
            futures = {pool.submit(_fund_category_asset_ids, category_id): category_id
                       for category_id in category_ids}
            for future in as_completed(futures):
                category_id = futures[future]
                try:
                    for asset_id in future.result():
                        mapping.setdefault(str(asset_id), category_id)
                except Exception:
                    continue
        FUND_CATEGORY_MAP_CACHE = (now, mapping)
        return mapping


def _schedule_fund_category_map_refresh() -> None:
    global FUND_CATEGORY_MAP_BUILDING
    now = time.monotonic()
    with FUND_CATEGORY_MAP_LOCK:
        if FUND_CATEGORY_MAP_BUILDING:
            return
        if _fund_category_map_is_fresh(FUND_CATEGORY_MAP_CACHE, now):
            return
        FUND_CATEGORY_MAP_BUILDING = True

    def refresh() -> None:
        global FUND_CATEGORY_MAP_BUILDING
        try:
            _fund_category_map()
        except Exception:
            pass
        finally:
            with FUND_CATEGORY_MAP_LOCK:
                FUND_CATEGORY_MAP_BUILDING = False

    Thread(target=refresh, name="rahavard-fund-category-map", daemon=True).start()


def _apply_fund_category_map(fund: dict[str, Any], category_map: dict[str, str]) -> None:
    raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
    asset_id = str(raw.get("rahavard_asset_id") or raw.get("asset_id")
                   or fund.get("registration_no") or fund.get("symbol") or "").strip()
    mapped = category_map.get(asset_id)
    if mapped and not _fund_category_id(fund):
        fund["category_id"] = mapped


def _exclude_fixed_income_funds(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    try:
        fixed_income_ids = _fund_category_asset_ids("255")
    except SourceError:
        fixed_income_ids = set()
    cached_category_map = FUND_CATEGORY_MAP_CACHE
    if cached_category_map:
        fixed_income_ids.update(asset_id for asset_id, category_id in cached_category_map[1].items()
                                if category_id == "255")

    visible: list[dict[str, Any]] = []
    for item in items:
        raw = item.get("raw") if isinstance(item.get("raw"), dict) else {}
        asset_id = str(raw.get("rahavard_asset_id") or item.get("registration_no") or item.get("symbol") or "").strip()
        group_ids = (item.get("category_id"), raw.get("groupId"), raw.get("group_id"),
                     raw.get("category_id"), raw.get("categoryId"))
        labels = (item.get("category"), item.get("name"), raw.get("category"),
                  raw.get("category_name"), raw.get("groupName"), raw.get("fund_type"))
        is_fixed_income = (
            asset_id in fixed_income_ids
            or any(str(value or "").strip() == "255" for value in group_ids)
            or any("درآمد ثابت" in str(value or "").casefold()
                   or "fixed income" in str(value or "").casefold() for value in labels)
        )
        if not is_fixed_income:
            visible.append(item)
    return visible


def _dedupe_rahavard_funds(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    selected: dict[str, tuple[int, dict[str, Any]]] = {}
    for index, item in enumerate(items):
        raw = item.get("raw") if isinstance(item.get("raw"), dict) else {}
        asset_id = str(raw.get("rahavard_asset_id") or item.get("registration_no") or item.get("symbol") or "").strip()
        identity = asset_id or str(item.get("fund_key") or index)
        current = selected.get(identity)
        if current is None or _fund_observation_rank(item) > _fund_observation_rank(current[1]):
            selected[identity] = (index, item)
    return [item for _, item in sorted(selected.values(), key=lambda entry: entry[0])]


@app.get("/api/funds")
def funds(search: str = "", category: str = "", sort: str = "data_timestamp",
          horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    if category == "255":
        return {"items": [], "count": 0, "source": "Rahavard365", "data_status": "available"}
    db_category = "" if category.isdigit() else category
    items = _dedupe_rahavard_funds(db.funds(search=search, category=db_category, sort=sort, source="Rahavard365"))
    if category.isdigit():
        local_matches = [item for item in items if _fund_category_id(item) == category]
        if local_matches:
            items = local_matches
        else:
            try:
                allowed_ids = _fund_category_asset_ids(category)
                items = [item for item in items
                         if str(item.get("raw", {}).get("rahavard_asset_id") or item.get("symbol") or "") in allowed_ids]
            except SourceError:
                items = []
    else:
        cached_map = FUND_CATEGORY_MAP_CACHE
        if _fund_category_map_is_fresh(cached_map):
            category_map = cached_map[1]
        else:
            category_map = {}
            _schedule_fund_category_map_refresh()
        for item in items:
            _apply_fund_category_map(item, category_map)
    items = _exclude_fixed_income_funds(items)
    for item in items:
        raw = item.get("raw") if isinstance(item.get("raw"), dict) else {}
        current_price = _watchlist_number(item.get("market_price"))
        if current_price is None or current_price <= 0:
            current_price = _watchlist_number(raw.get("real_close_price"))
        item["rahavard_asset_id"] = str(raw.get("rahavard_asset_id") or item["symbol"])
        item["rahavard_state"] = raw.get("instrument_state")
        item["updated_at"] = _fund_updated_at(item)
        item["price_is_cached"] = _timestamp_is_cached(item.get("data_timestamp") or item["updated_at"])
        for field in ("open_price", "high_price", "low_price", "ask_price", "ask_volume",
                      "bid_price", "bid_volume", "real_close_price_change",
                      "real_close_price_change_percent", "trade_date_time"):
            item[field] = raw.get(field)
        result = None
        category_id = _fund_category_id(item, category)
        model_context = _fund_model_context(
            item, category_id=category_id,
        )
        item["fund_type_key"] = model_context.get("resolved_fund_type")
        item["fund_type_label"] = FUND_TYPE_LABELS.get(item["fund_type_key"], item.get("category") or "نامشخص")
        item["analysis"] = result
        item["technical_score"] = result.get("technical_score") if result else None
        item["data_coverage_pct"] = result.get("data_coverage_pct") if result else None
        item["risk_score"] = (result.get("risk_rank") or {}).get("score") if result else None
        item["score_change_5d"] = result.get("score_change_5d") if result else None
        item["score_change_20d"] = result.get("score_change_20d") if result else None
        item["medium_term_exit_price"] = result.get("medium_term_exit_price") if result else None
        item["purchase_assessment"] = _fund_purchase_assessment(current_price, result)
        item["buy_suitability_rank"] = item["purchase_assessment"]["rank"]
        item["trend"] = result.get("trend") if result else "unavailable"
        item["indicator_fetched_at"] = None
        item["technical_source"] = "Rahavard365"
    sort_key = {"return": "one_year_return", "technical_score": "technical_score"}.get(sort, sort)
    if sort_key in {"one_year_return", "monthly_return", "daily_return", "technical_score"}:
        items.sort(key=lambda f: f.get(sort_key) if f.get(sort_key) is not None else float("-inf"), reverse=True)
    return {"items": items, "count": len(items), "source": "Rahavard365",
            "data_status": "available" if items else "unavailable"}


def _fund_or_error(identifier: str) -> dict[str, Any]:
    fund = db.get_fund(identifier)
    if fund and fund.get("ambiguous"):
        matches = [{"fund_key": row["fund_key"], "name": row["name"]} for row in fund["matches"]]
        raise HTTPException(status_code=409, detail={"message": "این شناسه به چند صندوق تعلق دارد؛ از fund_key استفاده کنید.", "matches": matches})
    if not fund:
        raise HTTPException(status_code=404, detail="صندوق پیدا نشد یا هنوز از منبع دریافت نشده است")
    if fund.get("source") != "Rahavard365":
        raise HTTPException(status_code=404, detail="این صندوق در فهرست ره‌آورد۳۶۵ موجود نیست")
    return fund


def _latest_fund_nav(fund_key: str) -> tuple[float | None, str | None]:
    for row in reversed(db.fund_nav_history(fund_key)):
        for key in ("statistical_nav", "cancel_nav", "issue_nav"):
            value = _watchlist_number(row.get(key))
            if value is not None and value > 0:
                stamp = row.get("timestamp")
                return value, str(stamp) if stamp is not None else None
    return None, None


@app.post("/api/funds/nav/batch")
async def fund_nav_batch(payload: FundNavBatchInput):
    keys = list(dict.fromkeys(key.strip() for key in payload.fund_keys if key.strip()))
    if not keys:
        raise HTTPException(status_code=422, detail="شناسه‌ای برای محاسبه NAV ارسال نشده است")

    semaphore = asyncio.Semaphore(4)

    async def calculate(identifier: str) -> dict[str, Any]:
        async with semaphore:
            try:
                fund = _fund_or_error(identifier)
                nav, nav_timestamp = _latest_fund_nav(fund["fund_key"])
                if nav is None or stale_status(nav_timestamp, max_hours=MAX_STALE_HOURS):
                    try:
                        await fetch_rahavard_fund_profile(fund["fund_key"])
                    except Exception:
                        pass
                    nav, nav_timestamp = _latest_fund_nav(fund["fund_key"])

                raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
                price = _watchlist_number(fund.get("market_price"))
                if price is None or price <= 0:
                    price = _watchlist_number(raw.get("real_close_price"))
                premium = ((price / nav - 1) * 100
                           if price is not None and price > 0 and nav is not None and nav > 0
                           else None)
                return {"fund_key": fund["fund_key"], "nav": nav,
                        "nav_timestamp": nav_timestamp,
                        "nav_is_cached": stale_status(nav_timestamp, max_hours=MAX_STALE_HOURS) is not False,
                        "nav_premium_pct": round(premium, 3) if premium is not None else None}
            except HTTPException as exc:
                return {"fund_key": identifier, "nav": None, "nav_premium_pct": None,
                        "error": str(exc.detail)}
            except Exception as exc:
                return {"fund_key": identifier, "nav": None, "nav_premium_pct": None,
                        "error": str(exc)[:180]}

    items = await asyncio.gather(*(calculate(key) for key in keys))
    return {"items": items, "processed": len(items),
            "valued": sum(item.get("nav_premium_pct") is not None for item in items)}


@app.post("/api/funds/scores/batch")
async def fund_score_batch(payload: FundScoreBatchInput,
                           horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    keys = list(dict.fromkeys(key.strip() for key in payload.fund_keys if key.strip()))
    if not keys:
        raise HTTPException(status_code=422, detail="شناسه‌ای برای محاسبه امتیاز ارسال نشده است")

    category_id = payload.category_id if payload.category_id in FUND_TYPE_BY_CATEGORY_ID else ""
    peer_groups: dict[str, list[dict[str, Any]]] = {}
    if category_id:
        try:
            peer_asset_ids = _fund_category_asset_ids(category_id)
            peers = [item for item in _dedupe_rahavard_funds(
                db.funds(source="Rahavard365"))
                if str((item.get("raw") or {}).get("rahavard_asset_id") or item.get("symbol") or "") in peer_asset_ids]
            peer_groups[category_id] = peers
        except SourceError:
            pass
    else:
        category_map = await asyncio.to_thread(_fund_category_map)
        for peer in db.funds(source="Rahavard365"):
            _apply_fund_category_map(peer, category_map)
            peer_key = _fund_peer_group_key(peer)
            if peer_key:
                peer_groups.setdefault(peer_key, []).append(peer)

    semaphore = asyncio.Semaphore(4)

    async def calculate(identifier: str) -> dict[str, Any]:
        async with semaphore:
            try:
                fund = _fund_or_error(identifier)
                if not category_id:
                    _apply_fund_category_map(fund, category_map)
                client_bars = None
                profile = None
                if payload.chart_bars is not None:
                    client_bars = payload.chart_bars.get(fund["fund_key"])
                    if not client_bars:
                        raise SourceError("کندل نمودار ره‌آورد برای این نماد دریافت نشد")
                raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
                current_price = _watchlist_number(fund.get("market_price"))
                if current_price is None or current_price <= 0:
                    current_price = _watchlist_number(raw.get("real_close_price"))
                actual_category_id = _fund_category_id(fund, category_id)
                peer_key = actual_category_id or _fund_peer_group_key(fund)
                model_context = _fund_model_context(
                    fund, category_id=actual_category_id,
                    peer_items=peer_groups.get(peer_key or ""))
                analysis_result, chart_bars, _ = await _analyze_fund_chart(
                    fund, profile=profile, model_context=model_context, horizon=horizon,
                    preloaded_bars=client_bars)
                _record_fund_score_history(analysis_result, fund["fund_key"],
                                           analysis_result.get("data_timestamp"), chart_bars)
                purchase_assessment = _fund_purchase_assessment(current_price, analysis_result)
                return {"fund_key": fund["fund_key"], "technical_score": analysis_result.get("technical_score"),
                        "technical_source": "Rahavard365",
                        "analysis": analysis_result,
                        "investment_horizon": analysis_result.get("investment_horizon"),
                        "investment_horizon_label": analysis_result.get("investment_horizon_label"),
                        "medium_term_exit_price": analysis_result.get("medium_term_exit_price"),
                        "purchase_assessment": purchase_assessment,
                        "buy_suitability_rank": purchase_assessment["rank"],
                        "data_coverage_pct": analysis_result.get("data_coverage_pct"),
                        "risk_score": (analysis_result.get("risk_rank") or {}).get("score"),
                        "score_change_5d": analysis_result.get("score_change_5d"),
                        "score_change_20d": analysis_result.get("score_change_20d"),
                        "fund_type_key": analysis_result.get("fund_type"),
                        "fund_type_label": analysis_result.get("fund_type_label"),
                        "error": None if analysis_result.get("technical_score") is not None else "داده کافی برای محاسبه موجود نیست"}
            except HTTPException as exc:
                return {"fund_key": identifier, "technical_score": None, "error": str(exc.detail)}
            except SourceError as exc:
                return {"fund_key": identifier, "technical_score": None, "error": str(exc)}
            except Exception as exc:
                return {"fund_key": identifier, "technical_score": None, "error": str(exc)[:180]}

    items = await asyncio.gather(*(calculate(key) for key in keys))
    return {"items": items, "processed": len(items),
            "scored": sum(item["technical_score"] is not None for item in items),
            "unavailable": sum(item["technical_score"] is None for item in items)}


def _watchlist_number(value: Any) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError, OverflowError):
        return None


def _latest_confirmed_daily_bar(bars: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Ignore today's still-forming candle until the Tehran session has closed."""
    tehran = timezone(timedelta(hours=3, minutes=30))
    now = datetime.now(tehran)
    session_close = datetime_time(13, 0)
    for bar in reversed(bars or []):
        stamp = bar.get("timestamp")
        if not stamp:
            continue
        try:
            dt = datetime.fromisoformat(str(stamp).replace("Z", "+00:00"))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            local_dt = dt.astimezone(tehran)
        except (TypeError, ValueError, OverflowError):
            continue
        if local_dt.date() < now.date() or (local_dt.date() == now.date() and now.time() >= session_close):
            return bar
    return None


def _fund_purchase_assessment(price: Any, analysis: dict[str, Any] | None) -> dict[str, Any]:
    current_price = _watchlist_number(price)
    if not isinstance(analysis, dict):
        return {"status": "unknown", "label": "داده کافی نیست", "reason": "تحلیل تکنیکال این نماد هنوز آماده نیست.", "rank": None}
    score_horizon = analysis.get("investment_horizon_label") or horizon_label(analysis.get("investment_horizon"))
    short_horizon = analysis.get("investment_horizon") == "short"
    exit_price = _watchlist_number(
        analysis.get("short_term_exit_price" if short_horizon else "medium_term_exit_price"))
    exit_average = "EMA(10)" if short_horizon else "SMA(30)"
    exit_horizon = "کوتاه‌مدت" if short_horizon else "میان‌مدت"
    score = _watchlist_number(analysis.get("technical_score"))
    breakdown = analysis.get("score_breakdown") if isinstance(analysis.get("score_breakdown"), dict) else {}
    family_scores = analysis.get("family_scores") if isinstance(analysis.get("family_scores"), dict) else {}
    moving_family = family_scores.get("moving_averages") if isinstance(family_scores.get("moving_averages"), dict) else {}
    moving_average_score = _watchlist_number(moving_family.get("score"))
    rsi = _watchlist_number(breakdown.get("rsi"))
    resistance_distance_pct = _watchlist_number(breakdown.get("nearest_resistance_distance_pct"))
    coverage = _watchlist_number(analysis.get("data_coverage_pct"))
    available_weight = _watchlist_number(analysis.get("available_family_weight_pct"))
    risk_level = (analysis.get("risk_rank") or {}).get("level")
    if current_price is None or current_price <= 0 or exit_price is None or exit_price <= 0 or score is None:
        return {"status": "unknown", "label": "داده کافی نیست", "reason": "برای ارزیابی قیمت فعلی، قیمت و دادهٔ کامل تکنیکال لازم است.", "rank": None}
    if current_price <= exit_price:
        return {"status": "wait", "label": f"صبر · زیر {exit_average}",
                "reason": f"قیمت فعلی زیر میانگین {exit_average} است؛ برای تشخیص شکست حد خروج {exit_horizon}، بسته‌شدن کندل روزانه را بررسی کن.", "rank": 2}
    if ((coverage is not None and coverage < 30)
            or (available_weight is not None and available_weight < 50)):
        return {"status": "wait", "label": "صبر · پوشش داده کم",
                "reason": "کمتر از نیمی از وزن مدل با دادهٔ معتبر پوشش دارد؛ امتیاز برای تصمیم خرید کافی نیست.", "rank": 2}
    if risk_level == "high":
        return {"status": "wait", "label": "صبر · ریسک بالا",
                "reason": "رتبهٔ ریسک نوسان بالاست؛ امتیاز تکنیکال به‌تنهایی برای ورود کافی نیست.", "rank": 2}
    decision = entry_recommendation(
        rsi, moving_average_score, score,
        above_exit_average=current_price > exit_price,
        resistance_distance_pct=resistance_distance_pct)
    status = decision["status"]
    if status == "suitable":
        rank, outward_status = 3, "suitable"
    elif status == "unsuitable":
        rank, outward_status = 1, "unsuitable"
    elif status == "unavailable":
        rank, outward_status = None, "unknown"
    else:
        rank, outward_status = 2, "wait"
    return {**decision, "status": outward_status, "rank": rank,
            "rsi": rsi, "moving_average_score": moving_average_score,
            "overall_score": score, "horizon": score_horizon}


def _add_position_from_assessment(assessment: dict[str, Any]) -> dict[str, str]:
    if assessment.get("status") == "suitable":
        return {"label": "بله · پله‌ای", "tone": "positive", "reason": assessment.get("reason", "شرایط RSI و میانگین‌ها مناسب است؛ مرحله‌ای اضافه کن.")}
    if assessment.get("status") == "unsuitable":
        return {"label": "خیر · فعلاً", "tone": "negative", "reason": assessment.get("reason", "شرایط ورود مناسب نیست.")}
    return {"label": assessment.get("label", "صبر · داده کافی نیست"),
            "tone": "neutral", "reason": assessment.get("reason", "برای افزایش حجم تا کامل‌شدن شرایط ورود صبر کن.")}


def _watchlist_match_key(value: Any) -> str:
    return str(value or "").strip().casefold().replace("ي", "ی").replace("ك", "ک")


def _watchlist_asset_id(fund: dict[str, Any]) -> str:
    raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
    return str(raw.get("rahavard_asset_id") or fund.get("registration_no") or fund.get("symbol") or "").strip()


def _watchlist_find_fund(identifier: str) -> dict[str, Any] | None:
    fund = db.get_fund(identifier)
    if fund and not fund.get("ambiguous") and fund.get("source") == "Rahavard365":
        return fund
    needle = _watchlist_match_key(identifier)
    matches = []
    for candidate in db.funds(source="Rahavard365"):
        raw = candidate.get("raw") if isinstance(candidate.get("raw"), dict) else {}
        exact_values = {
            candidate.get("fund_key"), candidate.get("symbol"), candidate.get("registration_no"),
            candidate.get("name"), raw.get("rahavard_asset_id"), raw.get("trade_symbol"),
        }
        if any(_watchlist_match_key(value) == needle for value in exact_values if value is not None):
            matches.append(candidate)
    unique: dict[str, dict[str, Any]] = {}
    for item in matches:
        asset_id = _watchlist_asset_id(item)
        identity = asset_id or item["fund_key"]
        previous = unique.get(identity)
        if previous is None or item["fund_key"] == f"rahavard:{asset_id}":
            unique[identity] = item
    if len(unique) > 1:
        raise HTTPException(status_code=409, detail="چند صندوق با این نام پیدا شد؛ شناسه ره‌آورد را وارد کن")
    return next(iter(unique.values()), None)


_DEFAULT_WATCHLIST_ITEMS = [
    ("آلتون", 7000, 75197),
    ("دارونو", 3500, 34822),
    ("سینرژی", 1000, 63651),
    ("آگاس", 100, 460104),
    ("متال", 2000, 40915),
    ("اهرم", 1500, 76341),
]
_LEGACY_WATCHLIST_SYMBOLS = {"سبزآبنوس", "تکپاد"}
_WATCHLIST_DISPLAY_BY_ASSET_ID = {
    "22820": "آلتون",
    "35665": "دارونو",
    "35129": "سینرژی",
    "820": "آگاس",
    "25249": "متال",
    "18335": "اهرم",
}
_DEFAULT_WATCHLIST_ASSET_ID_BY_SYMBOL = {
    symbol: asset_id for asset_id, symbol in _WATCHLIST_DISPLAY_BY_ASSET_ID.items()
}


def _ensure_default_watchlist_fund(symbol: str) -> dict[str, Any] | None:
    """Keep known default symbols resolvable when Rahavard omits them intraday."""
    asset_id = _DEFAULT_WATCHLIST_ASSET_ID_BY_SYMBOL.get(symbol)
    if not asset_id:
        return None
    fund_key = f"rahavard:{asset_id}"
    db.save_funds([{
        "fund_key": fund_key, "symbol": asset_id, "registration_no": asset_id,
        "name": symbol, "category": "قابل معامله", "category_id": "rahavard-etf",
        "market_price": None, "raw": {"rahavard_asset_id": asset_id, "trade_symbol": symbol},
    }], "Rahavard365", f"https://rahavard365.com/asset/{asset_id}")
    return db.get_fund(fund_key)



def _seed_default_watchlist() -> None:
    records = db.watchlist_items()
    desired_keys: set[str] = set()
    for symbol, units, break_even_price in _DEFAULT_WATCHLIST_ITEMS:
        try:
            fund = _watchlist_find_fund(symbol)
        except HTTPException:
            fund = None
        if fund is None:
            fund = _ensure_default_watchlist_fund(symbol)
        target_key = fund["fund_key"] if fund else f"manual:{symbol}"
        target_asset_id = _watchlist_asset_id(fund) if fund else None
        desired_keys.add(target_key)

        existing = None
        for record in records:
            if record["fund_key"] in {target_key, f"manual:{symbol}"}:
                existing = record
                break
            if target_asset_id:
                saved_fund = db.get_fund(record["fund_key"])
                if saved_fund and _watchlist_asset_id(saved_fund) == target_asset_id:
                    existing = record
                    break
        if existing and existing["fund_key"] != target_key:
            old_key = existing["fund_key"]
            record = db.add_watchlist_item(target_key)
            saved_units = existing.get("units")
            saved_break_even = existing.get("break_even_price")
            if symbol == "آلتون" and (saved_units, saved_break_even) == (6000, 73837):
                saved_units, saved_break_even = units, break_even_price
            if symbol == "اهرم":
                saved_units, saved_break_even = units, break_even_price
            if target_asset_id == "25249":
                saved_units, saved_break_even = units, break_even_price
            db.update_watchlist_item(target_key, {
                "units": saved_units,
                "break_even_price": saved_break_even,
            })
            db.remove_watchlist_item(old_key)
            existing = {**record, "fund_key": target_key}
            records = [row for row in records if row["fund_key"] != old_key]
            records.append(existing)
        elif existing:
            # Keep saved values exactly as entered, including an intentional blank.
            if symbol == "آلتون" and (existing.get("units"), existing.get("break_even_price")) == (6000, 73837):
                db.update_watchlist_item(target_key, {"units": units, "break_even_price": break_even_price})
            elif symbol == "اهرم" and (existing.get("units"), existing.get("break_even_price")) != (units, break_even_price):
                db.update_watchlist_item(target_key, {"units": units, "break_even_price": break_even_price})
            elif target_asset_id == "25249" and (existing.get("units"), existing.get("break_even_price")) != (units, break_even_price):
                db.update_watchlist_item(target_key, {"units": units, "break_even_price": break_even_price})
        else:
            record = db.add_watchlist_item(target_key)
            db.update_watchlist_item(target_key, {"units": units, "break_even_price": break_even_price})
            records.append({**record, "units": units, "break_even_price": break_even_price})

    # Remove the two rows that were seeded by the previous default portfolio.
    for record in list(records):
        if record["fund_key"] in desired_keys:
            continue
        saved_fund = db.get_fund(record["fund_key"])
        raw = saved_fund.get("raw") if isinstance(saved_fund, dict) else {}
        saved_symbols = {
            record["fund_key"].removeprefix("manual:"),
            saved_fund.get("symbol") if saved_fund else None,
            saved_fund.get("name") if saved_fund else None,
            raw.get("trade_symbol"),
        }
        if _LEGACY_WATCHLIST_SYMBOLS.intersection(saved_symbols):
            db.remove_watchlist_item(record["fund_key"])

def _watchlist_levels(indicator_data: dict[str, Any] | None,
                      price: float | None) -> dict[str, list[dict[str, Any]]]:
    if not indicator_data or price is None:
        return {"supports": [], "resistances": []}
    pivots = indicator_data.get("pivots") if isinstance(indicator_data.get("pivots"), list) else []
    supports: list[dict[str, Any]] = []
    resistances: list[dict[str, Any]] = []
    for item in pivots:
        if not isinstance(item, dict):
            continue
        method = str(item.get("short_name_en") or item.get("name_en") or "Pivot")
        parts = item.get("value") if isinstance(item.get("value"), list) else []
        for part in parts:
            if not isinstance(part, dict):
                continue
            label = str(part.get("name") or "").strip().casefold()
            if label not in {"r1", "r2", "r3", "r4", "s1", "s2", "s3", "s4"}:
                continue
            level = _watchlist_number(part.get("value"))
            if level is None or level <= 0:
                continue
            entry = {"price": level, "label": f"{method} · {label.upper()}"}
            (supports if level < price else resistances).append(entry)

    def nearest(values: list[dict[str, Any]], reverse: bool) -> list[dict[str, Any]]:
        values.sort(key=lambda row: row["price"], reverse=reverse)
        distinct: list[dict[str, Any]] = []
        for row in values:
            if any(abs(row["price"] - old["price"]) / max(old["price"], 1) < 0.001 for old in distinct):
                continue
            distinct.append(row)
            if len(distinct) == 3:
                break
        return distinct

    return {"supports": nearest(supports, True), "resistances": nearest(resistances, False)}


@app.get("/api/watchlist")
async def get_watchlist(horizon: str = Query(default="long"),
                        skip_technical: bool = Query(default=False)):
    horizon = normalize_horizon(horizon)
    _seed_default_watchlist()
    records = db.watchlist_items()
    semaphore = asyncio.Semaphore(5)

    async def load_item(record: dict[str, Any]) -> dict[str, Any]:
        fund = db.get_fund(record["fund_key"])
        if not fund or fund.get("source") != "Rahavard365":
            is_custom_symbol = str(record["fund_key"]).startswith("manual:")
            custom_symbol = str(record["fund_key"])[len("manual:"):] if is_custom_symbol else ""
            display_symbol = custom_symbol or record["fund_key"]
            return {**record, "name": display_symbol, "symbol": display_symbol,
                    "is_custom_symbol": is_custom_symbol, "price": None,
                    "supports": [], "resistances": [],
                    "midterm_outlook": "نماد سفارشی · دادهٔ قیمت موجود نیست" if is_custom_symbol else "صندوق در فهرست ره‌آورد پیدا نشد",
                    "overall_status": "داده در دسترس نیست", "technical_score": None,
                    "short_term_exit_price": None, "medium_term_exit_price": None,
                    "active_exit_price": None, "active_exit_horizon": horizon_label(horizon),
                    "active_exit_average": "EMA(10)" if horizon == "short" else "SMA(30)",
                    "distance_to_exit": None,
                    "holding_decision": {"label": "داده کافی نیست", "tone": "unknown", "reason": "قیمت و تحلیل تکنیکال این نماد در دسترس نیست."},
                    "add_position_decision": {"label": "فعلاً اضافه نکن", "tone": "neutral", "reason": "تا دریافت قیمت و تحلیل معتبر صبر کن."},
                    "explanation": "برای این نماد، نمایه و دادهٔ تحلیلی در فهرست فعلی صندوق‌های ره‌آورد موجود نیست."}
        async with semaphore:
            try:
                profile = await fetch_rahavard_fund_profile(record["fund_key"])
            except Exception:
                profile = {}
        asset = profile.get("asset") if isinstance(profile.get("asset"), dict) else {}
        trade = profile.get("last_trade") if isinstance(profile.get("last_trade"), dict) else {}
        raw = fund.get("raw") if isinstance(fund.get("raw"), dict) else {}
        live_price = _watchlist_number(trade.get("real_close_price"))
        if live_price is not None and live_price <= 0:
            live_price = None
        live_timestamp = trade.get("end_date_time")
        live_price_is_today = False
        if live_price is not None and live_timestamp:
            try:
                iran_tz = timezone(timedelta(hours=3, minutes=30))
                observed = datetime.fromisoformat(str(live_timestamp).replace("Z", "+00:00"))
                if observed.tzinfo is None:
                    observed = observed.replace(tzinfo=iran_tz)
                live_price_is_today = observed.astimezone(iran_tz).date() == datetime.now(iran_tz).date()
            except (TypeError, ValueError, OverflowError):
                live_price_is_today = False
        cached_history = db.history(record["fund_key"], 3, source="Rahavard365")
        cached_price = _watchlist_number(fund.get("market_price"))
        if len(cached_history) < 2:
            try:
                cached_history = await fetch_fund_price_history(record["fund_key"])
            except Exception:
                cached_history = db.history(record["fund_key"], 3, source="Rahavard365")
        if cached_price is None and cached_history:
            cached_price = _watchlist_number(cached_history[-1].get("close"))
        price = live_price or cached_price
        price_is_cached = price is not None and not live_price_is_today
        price_note = ("قیمت پایانی آخرین روز معاملاتی؛ با ثبت معامله جدید به‌روز می‌شود"
                      if live_price is not None and not live_price_is_today
                      else "آخرین قیمت ذخیره‌شده قبل از شروع معاملات"
                      if live_price is None and cached_price is not None
                      else "قیمت جاری ره‌آورد۳۶۵")
        change_ratio = _watchlist_number(trade.get("real_close_price_change_percent"))
        daily_return = change_ratio * 100 if change_ratio is not None else fund.get("daily_return")
        analysis_result = None
        chart_bars: list[dict[str, Any]] = []
        indicator_data: dict[str, Any] | None = None
        if not skip_technical:
            try:
                analysis_result, chart_bars, indicator_data = await _analyze_fund_chart(
                    fund, profile=profile, horizon=horizon)
            except Exception:
                pass
        confirmed_daily_bar = _latest_confirmed_daily_bar(chart_bars)
        last_daily_close = _watchlist_number(confirmed_daily_bar.get("close")) if confirmed_daily_bar else None
        stamp = (trade.get("end_date_time") if live_price is not None else None) \
            or ((cached_history[-1].get("timestamp") if cached_history else None)) \
            or fund.get("data_timestamp")
        technical_price = _watchlist_number(analysis_result.get("price")) if analysis_result else None
        levels = _watchlist_levels(indicator_data, technical_price or price)
        signal = analysis_result.get("signal") if analysis_result else None
        trend = analysis_result.get("trend") if analysis_result else None
        short_term_exit_price = _watchlist_number(analysis_result.get("short_term_exit_price")) if analysis_result else None
        medium_term_exit_price = _watchlist_number(analysis_result.get("medium_term_exit_price")) if analysis_result else None
        active_exit_price = short_term_exit_price if horizon == "short" else medium_term_exit_price
        active_exit_horizon = horizon_label(horizon)
        active_exit_average = "EMA(10)" if horizon == "short" else "SMA(30)"
        exit_confirmed = (last_daily_close is not None and active_exit_price is not None
                          and last_daily_close < active_exit_price)
        distance_value = price - active_exit_price if price is not None and active_exit_price is not None else None
        distance_pct = distance_value / price * 100 if distance_value is not None and price else None
        break_even = _watchlist_number(record.get("break_even_price"))
        units = _watchlist_number(record.get("units"))
        profit_per_unit = price - break_even if price is not None and break_even is not None else None
        profit_pct = profit_per_unit / break_even * 100 if profit_per_unit is not None and break_even else None
        profit_total = profit_per_unit * units if profit_per_unit is not None and units is not None else None
        if signal == "buy":
            outlook, tone = "تمایل مثبت · نگهداری تحت نظر", "positive"
        elif signal == "sell":
            outlook, tone = "تمایل منفی · بررسی فروش", "negative"
        elif signal == "hold":
            outlook, tone = "خنثی · نیازمند بررسی بیشتر", "neutral"
        else:
            outlook, tone = "داده تکنیکال کافی نیست", "unknown"
        signal_fa = {"buy": "مثبت", "sell": "منفی", "hold": "خنثی"}.get(signal, "نامشخص")
        technical_score = _watchlist_number(analysis_result.get("technical_score")) if analysis_result else None
        score_horizon = analysis_result.get("investment_horizon_label") if analysis_result else horizon_label(horizon)
        score_breakdown = analysis_result.get("score_breakdown") if analysis_result and isinstance(analysis_result.get("score_breakdown"), dict) else {}
        purchase_assessment = _fund_purchase_assessment(price, analysis_result)
        if price is None or active_exit_price is None or technical_score is None:
            holding_decision = {"label": "داده کافی نیست", "tone": "unknown", "reason": "برای تعیین نگهداری یا خروج، قیمت و تحلیل کامل لازم است."}
            add_position_decision = {"label": "فعلاً اضافه نکن", "tone": "neutral", "reason": "تا کامل‌شدن داده‌ها برای افزایش حجم صبر کن."}
        elif last_daily_close is None:
            holding_decision = ({"label": "موقعیتی ثبت نشده", "tone": "neutral", "reason": "برای این نماد تعداد واحدی در دیده‌بان ثبت نشده است."}
                                if (units or 0) <= 0 else
                                {"label": "در انتظار کندل روزانه", "tone": "neutral", "reason": f"آخرین قیمت پایانی روزانه برای تأیید قاعدهٔ SMA در دسترس نیست؛ قیمت لحظه‌ای به‌تنهایی مبنای خروج نیست."})
            add_position_decision = {"label": "صبر تا تأیید روزانه", "tone": "neutral", "reason": "تا دریافت قیمت پایانی روزانه، شکست SMA و شرایط افزودن حجم را تأیید نمی‌کنم."}
        elif exit_confirmed:
            holding_decision = ({"label": "موقعیتی ثبت نشده", "tone": "neutral", "reason": "برای این نماد تعداد واحدی در دیده‌بان ثبت نشده است."}
                                if (units or 0) <= 0 else
                                {"label": "فروش / کاهش", "tone": "negative", "reason": f"آخرین کندل روزانه زیر {active_exit_average} بسته شده و حد خروج {active_exit_horizon} تأیید شده است."})
            add_position_decision = {"label": "خیر · حد خروج تأیید شد", "tone": "negative", "reason": f"تا زمانی که کندل روزانه بالای {active_exit_average} تثبیت نشده، حجم اضافه نکن."}
        elif (units or 0) <= 0:
            holding_decision = {"label": "موقعیتی ثبت نشده", "tone": "neutral", "reason": "برای این نماد تعداد واحدی در دیده‌بان ثبت نشده است."}
            add_position_decision = _add_position_from_assessment(purchase_assessment)
        elif technical_score < 40:
            holding_decision = {"label": "کاهش ریسک", "tone": "negative", "reason": f"امتیاز تکنیکال {score_horizon} ضعیف است؛ نگهداری کامل ریسک بالاتری دارد."}
            add_position_decision = {"label": "خیر · فعلاً", "tone": "negative", "reason": f"امتیاز {score_horizon} ضعیف است."}
        else:
            holding_decision = {"label": "نگهداری", "tone": "positive", "reason": f"آخرین کندل روزانه بالای {active_exit_average} بسته شده؛ حد خروج فقط با بسته‌شدن روزانه زیر آن فعال می‌شود."}
            add_position_decision = _add_position_from_assessment(purchase_assessment)
        asset_id = str(raw.get("rahavard_asset_id") or fund.get("registration_no") or fund["symbol"])
        display_symbol = _WATCHLIST_DISPLAY_BY_ASSET_ID.get(asset_id) or asset.get("trade_symbol") or fund["symbol"]
        return {
            **record,
            "name": display_symbol,
            "symbol": display_symbol,
            "fund_name": asset.get("name") or fund["name"],
            "trade_symbol": asset.get("trade_symbol") or raw.get("trade_symbol") or raw.get("name") or fund["name"],
            "rahavard_asset_id": asset_id,
            "price": price,
            "price_is_cached": price_is_cached,
            "price_note": price_note,
            "daily_return": daily_return,
            "data_timestamp": stamp,
            "supports": levels["supports"],
            "resistances": levels["resistances"],
            "profit": {"per_unit": profit_per_unit, "pct": profit_pct, "total": profit_total},
            "signal": signal,
            "signal_fa": signal_fa,
            "trend": trend,
            "technical_score": technical_score,
            "investment_horizon": analysis_result.get("investment_horizon") if analysis_result else horizon,
            "investment_horizon_label": score_horizon,
            "short_term_exit_price": short_term_exit_price,
            "medium_term_exit_price": medium_term_exit_price,
            "active_exit_price": active_exit_price,
            "active_exit_horizon": active_exit_horizon,
            "active_exit_average": active_exit_average,
            "last_daily_close": last_daily_close,
            "last_daily_close_timestamp": confirmed_daily_bar.get("timestamp") if confirmed_daily_bar else None,
            "technical_data_timestamp": analysis_result.get("data_timestamp") if analysis_result else None,
            "distance_to_exit": {"value": distance_value, "pct": distance_pct} if distance_value is not None else None,
            "holding_decision": holding_decision,
            "add_position_decision": add_position_decision,
            "purchase_assessment": purchase_assessment,
            "midterm_outlook": outlook,
            "outlook_tone": tone,
            "overall_status": analysis_result.get("explanation") if analysis_result else "دادهٔ نمودار ره‌آورد برای تحلیل در دسترس نیست.",
            "source": "Rahavard365",
            "technical_source": "Rahavard365",
        }

    items = await asyncio.gather(*(load_item(record) for record in records))
    default_order = {symbol: index for index, (symbol, _units, _price) in enumerate(_DEFAULT_WATCHLIST_ITEMS)}
    items.sort(key=lambda item: (default_order.get(item.get("symbol"), len(default_order)), str(item.get("fund_key", ""))))
    total_units = sum(_watchlist_number(item.get("units")) or 0 for item in items)
    positions = [item for item in items if (_watchlist_number(item.get("units")) or 0) > 0]
    has_cost = bool(positions) and all(_watchlist_number(item.get("break_even_price")) is not None for item in positions)
    has_market_value = bool(positions) and all(_watchlist_number(item.get("price")) is not None for item in positions)
    cost_basis = (sum(_watchlist_number(item.get("units")) * _watchlist_number(item.get("break_even_price"))
                       for item in positions) if has_cost else None)
    market_value = (sum(_watchlist_number(item.get("units")) * _watchlist_number(item.get("price"))
                        for item in positions) if has_market_value else None)
    pnl = market_value - cost_basis if has_cost and has_market_value else None
    totals = {
        "units": total_units,
        "cost_basis": cost_basis,
        "market_value": market_value,
        "pnl": pnl,
        "pnl_pct": (pnl / cost_basis * 100) if pnl is not None and cost_basis else None,
    }
    return {"items": items, "count": len(records), "source": "Rahavard365", "totals": totals,
            "storage": db.watchlist_storage_status()}


@app.post("/api/watchlist")
def add_watchlist_item(data: WatchlistAddInput):
    fund = _watchlist_find_fund(data.symbol)
    if not fund or fund.get("source") != "Rahavard365":
        raise HTTPException(status_code=404, detail="صندوقی با این شناسه در فهرست ره‌آورد پیدا نشد")
    record = None
    for existing in db.watchlist_items():
        saved_fund = db.get_fund(existing["fund_key"])
        if saved_fund and _watchlist_asset_id(saved_fund) == _watchlist_asset_id(fund):
            record = existing
            break
    existed = record is not None
    if record is None:
        record = db.add_watchlist_item(fund["fund_key"])
    return {"item": record, "fund": {"name": fund["name"], "symbol": fund["symbol"]},
            "created": not existed}


@app.get("/api/funds/{symbol}")
async def fund_detail(symbol: str, horizon: str = Query(default="long"),
                      skip_technical: bool = Query(default=False)):
    horizon = normalize_horizon(horizon)
    fund = _fund_or_error(symbol)
    fund_key = fund["fund_key"]
    profile_result, bars_result = await asyncio.gather(
        fetch_rahavard_fund_profile(fund_key),
        fetch_fund_price_history(fund_key),
        return_exceptions=True,
    )
    profile = profile_result if isinstance(profile_result, dict) else None
    bars = bars_result if isinstance(bars_result, list) else db.history(fund_key, 5000, source="Rahavard365")
    nav = db.fund_nav_history(fund_key)
    fund = dict(fund)
    fund["rahavard_asset_id"] = (fund.get("raw") or {}).get("rahavard_asset_id") or fund["registration_no"]
    profile_category = None
    latest_nav_timestamp = None
    if profile:
        asset = profile.get("asset") or {}
        spec = profile.get("fund_specification") or {}
        category = asset.get("category") or {}
        profile_category = category
        last_trade = profile.get("last_trade") or {}
        values = profile.get("fund_values") or []
        latest_nav = next((row for row in values if row.get("date")), {})
        latest_nav_timestamp = latest_nav.get("date") or latest_nav.get("timestamp")
        positive_navs = [row for row in values if (row.get("nav") or 0) > 0]
        aum_row = positive_navs[0] if positive_navs else latest_nav
        display_name = lambda value: value.get("name") if isinstance(value, dict) else value
        def number(value):
            try:
                parsed = float(value)
                return parsed if parsed == parsed and abs(parsed) != float("inf") else None
            except (TypeError, ValueError):
                return None
        fund.update({
            "name": asset.get("name") or fund["name"],
            "display_symbol": asset.get("trade_symbol") or fund["symbol"],
            "trade_symbol": asset.get("trade_symbol") or fund["symbol"],
            "category": display_name(category) or fund.get("category"),
            "manager": display_name(spec.get("manager")),
            "custodian": display_name(spec.get("custodian")),
            "market_maker": display_name(spec.get("market_maker")),
            "start_date": spec.get("inception_date"),
            "market_price": number(last_trade.get("real_close_price")) or fund.get("market_price"),
            "volume": number(last_trade.get("volume")) or fund.get("volume"),
            "value": number(last_trade.get("value")) or fund.get("value"),
            "trades": number(last_trade.get("trade_count")),
            "daily_return": (number(last_trade.get("real_close_price_change_percent")) * 100
                             if number(last_trade.get("real_close_price_change_percent")) is not None
                             else fund.get("daily_return")),
            "data_timestamp": last_trade.get("end_date_time") or fund.get("data_timestamp"),
        })
        fund["nav"] = next((number(latest_nav.get(key)) for key in ("statistical_price", "redemption_price", "bid_price")
                            if number(latest_nav.get(key)) is not None), fund.get("nav"))
        fund["aum"] = number(aum_row.get("nav")) or fund.get("aum")
        if fund.get("market_price") and fund.get("nav"):
            fund["nav_premium_pct"] = (fund["market_price"] / fund["nav"] - 1) * 100
    fund["price_is_cached"] = False
    if fund.get("market_price") is not None and fund.get("data_timestamp"):
        try:
            iran_tz = timezone(timedelta(hours=3, minutes=30))
            observed = datetime.fromisoformat(str(fund["data_timestamp"]).replace("Z", "+00:00"))
            if observed.tzinfo is None:
                observed = observed.replace(tzinfo=iran_tz)
            fund["price_is_cached"] = observed.astimezone(iran_tz).date() != datetime.now(iran_tz).date()
        except (TypeError, ValueError, OverflowError):
            fund["price_is_cached"] = False
    if not fund.get("market_price") and bars:
        fund["market_price"] = bars[-1].get("close")
        fund["data_timestamp"] = bars[-1].get("timestamp") or fund.get("data_timestamp")
        fund["price_is_cached"] = fund["market_price"] is not None
    if fund.get("price_is_cached"):
        fund["price_note"] = "آخرین قیمت پایانی ذخیره‌شده؛ با ثبت معامله جدید به‌روز می‌شود"
    model_context = _fund_model_context(
        fund, category=profile_category, price_history=[],
        nav_timestamp=latest_nav_timestamp)
    analysis_result = None
    technical_bars: list[dict[str, Any]] = []
    technical_source_error = None
    if not skip_technical:
        try:
            analysis_result, technical_bars, _ = await _analyze_fund_chart(
                fund, horizon=horizon, profile=profile, model_context=model_context,
                preloaded_bars=bars)
        except SourceError as exc:
            technical_source_error = str(exc)
    if analysis_result:
        _record_fund_score_history(analysis_result, fund_key,
                                   analysis_result.get("data_timestamp"), technical_bars)
    return {"fund": fund, "history": bars, "nav_history": nav, "portfolio": [],
            "analysis": analysis_result,
            "technical_source_error": technical_source_error,
            "purchase_assessment": _fund_purchase_assessment(fund.get("market_price"), analysis_result),
            "market_data_available": bool(bars),
            "fundamentals_source": "Rahavard365", "market_price_source": "Rahavard365" if fund.get("market_price") is not None else None,
            "history_source": "Rahavard365" if bars else None,
            "technical_source": "Rahavard365"}


@app.get("/api/funds/{symbol}/history")
async def fund_history(symbol: str):
    try:
        nav, price = await asyncio.gather(fetch_fund_nav_history(symbol), fetch_fund_price_history(symbol))
        return {"symbol": symbol, "nav_history": nav, "price_history": price,
                "sources": {"nav": "Rahavard365", "price": "Rahavard365"},
                "data_timestamp": {"nav": nav[-1]["timestamp"] if nav else None,
                                   "price": price[-1]["timestamp"] if price else None}}
    except SourceError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.get("/api/funds/{symbol}/portfolio")
async def fund_portfolio(symbol: str):
    details = await fund_detail(symbol)
    return {"symbol": symbol, "portfolio": details["portfolio"], "source": "Rahavard365",
            "note": details.get("portfolio_note"),
            "data_timestamp": details["fund"].get("data_timestamp")}


@app.get("/api/funds/{symbol}/technical")
async def fund_technical(symbol: str, horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    fund = _fund_or_error(symbol)
    try:
        result, bars, _ = await _analyze_fund_chart(fund, horizon=horizon)
    except SourceError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    _record_fund_score_history(result, fund["fund_key"], result.get("data_timestamp"), bars)
    return result


@app.post("/api/funds/{symbol}/technical")
async def fund_technical_from_client(symbol: str, payload: FundTechnicalBarsInput,
                                     horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    fund = _fund_or_error(symbol)
    try:
        result, bars, _ = await _analyze_fund_chart(
            fund, horizon=horizon, preloaded_bars=payload.bars)
    except SourceError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    _record_fund_score_history(result, fund["fund_key"], result.get("data_timestamp"), bars)
    db.save_source("rahavard_fund_charts", "نمودار و محاسبهٔ تکنیکال صندوق / ره‌آورد۳۶۵", "connected",
                   bars[-1].get("source_url"), bars[-1].get("timestamp"), None,
                   payload.chart_fetch_ms)
    return {**result, "purchase_assessment": _fund_purchase_assessment(fund.get("market_price"), result)}


@app.get("/api/funds/{symbol}/signals")
async def fund_signals(symbol: str, horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    fund = _fund_or_error(symbol)
    try:
        result, bars, _ = await _analyze_fund_chart(fund, horizon=horizon)
    except SourceError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    _record_fund_score_history(result, fund["fund_key"], result.get("data_timestamp"), bars)
    return {"symbol": fund["symbol"], "fund_key": fund["fund_key"], "source": "Rahavard365",
            "signal": result["signal"], "score": result["technical_score"],
            "investment_horizon": result["investment_horizon"],
            "investment_horizon_label": result["investment_horizon_label"],
            "gauges": result["site_gauges"], "indicator_notes": result["indicator_notes"],
            "explanation": result["explanation"], "scenarios": result.get("scenarios", []),
            "scenario_note": result.get("scenario_note")}


@app.get("/api/funds/{symbol}/analysis")
async def fund_analysis(symbol: str, horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    fund = _fund_or_error(symbol)
    try:
        result, bars, _ = await _analyze_fund_chart(fund, horizon=horizon)
    except SourceError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    _record_fund_score_history(result, fund["fund_key"], result.get("data_timestamp"), bars)
    return result


@app.get("/api/admin")
def admin():
    health_data = health()
    return {"health": health_data, "sources": db.source_status(), "recent_logs": db.fetch_logs(),
            "scheduler": REFRESH_STATE, "database_path": str(db.path)}


class AlertInput(BaseModel):
    symbol: str = Field(min_length=1, max_length=80)
    condition: str = Field(min_length=1, max_length=80)
    threshold: float | None = None


@app.get("/api/alerts")
def alerts():
    items = db.alerts()
    for item in items:
        fund = db.get_fund(item["symbol"])
        if fund and not fund.get("ambiguous"):
            item["display_symbol"] = fund["symbol"]
            item["fund_name"] = fund["name"]
        else:
            item["display_symbol"] = item["symbol"]
    return {"items": items}


@app.post("/api/alerts")
def create_alert(data: AlertInput):
    allowed = {"price_above", "price_below", "rsi_overbought", "rsi_oversold", "volume_spike"}
    if data.condition not in allowed:
        raise HTTPException(status_code=422, detail=f"شرط مجاز: {', '.join(sorted(allowed))}")
    if data.condition in {"price_above", "price_below"} and data.threshold is None:
        raise HTTPException(status_code=422, detail="برای شرط قیمت باید آستانه وارد شود")
    fund = db.get_fund(data.symbol)
    if fund and fund.get("ambiguous"):
        matches = [{"fund_key": row["fund_key"], "name": row["name"]} for row in fund["matches"]]
        raise HTTPException(status_code=409, detail={"message": "نماد تکراری است؛ برای هشدار صندوق از fund_key استفاده کنید.", "matches": matches})
    if not fund:
        with db.connect() as conn:
            known_asset = conn.execute("SELECT 1 FROM assets WHERE symbol=?", (data.symbol,)).fetchone()
        if not known_asset:
            raise HTTPException(status_code=404, detail="نماد یا شناسه صندوق شناخته‌شده نیست")
    alert_id = db.add_alert(data.symbol, data.condition, data.threshold)
    return {"id": alert_id, "status": "created"}


@app.get("/api/settings")
def settings():
    saved = db.get_settings()
    return {"data_update_interval": _refresh_interval(),
            "max_stale_hours": MAX_STALE_HOURS,
            "theme": saved.get("theme", "dark"),
            "language": "fa-IR", "settings": saved}


class SettingInput(BaseModel):
    key: str
    value: str


@app.put("/api/settings")
def update_setting(data: SettingInput):
    allowed = {"theme", "refresh_interval_seconds"}
    if data.key not in allowed:
        raise HTTPException(status_code=422, detail="این تنظیم از رابط کاربری قابل تغییر نیست")
    if data.key == "theme" and data.value not in {"dark", "light"}:
        raise HTTPException(status_code=422, detail="مقدار تم معتبر نیست")
    if data.key == "refresh_interval_seconds":
        try:
            interval = int(data.value)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="فاصله به‌روزرسانی باید عدد باشد") from exc
        if not 60 <= interval <= 86400:
            raise HTTPException(status_code=422, detail="بازه مجاز ۶۰ تا ۸۶۴۰۰ ثانیه است")
    db.save_setting(data.key, data.value)
    if data.key == "refresh_interval_seconds" and REFRESH_WAKE is not None:
        REFRESH_WAKE.set()
    return {"saved": True, "key": data.key, "value": data.value}


@app.get("/api/export/funds.csv")
def export_funds_csv():
    items = db.funds(source="Rahavard365")
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=["symbol", "name", "category", "manager", "aum", "nav", "market_price", "nav_premium_pct", "volume", "value", "data_timestamp"], extrasaction="ignore")
    writer.writeheader()
    writer.writerows(items)
    return StreamingResponse(iter(["\ufeff" + stream.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": "attachment; filename=funds.csv"})


@app.get("/api/export/funds.xlsx")
def export_funds_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise HTTPException(status_code=501, detail="برای خروجی Excel، openpyxl را نصب کنید") from exc
    items = db.funds(source="Rahavard365")
    columns = [("symbol", "نماد"), ("name", "نام"), ("category", "دسته"), ("manager", "مدیر"),
               ("daily_return", "بازده روزانه"), ("weekly_return", "بازده هفتگی"),
               ("monthly_return", "بازده ماهانه"), ("three_month_return", "بازده ۳ماهه"),
               ("six_month_return", "بازده ۶ماهه"), ("one_year_return", "بازده یک‌ساله"),
               ("nav", "NAV"), ("market_price", "قیمت بازار"), ("nav_premium_pct", "فاصله NAV٪"),
               ("aum", "دارایی تحت مدیریت"), ("volume", "حجم"), ("value", "ارزش"), ("data_timestamp", "زمان داده")]
    wb = Workbook(); ws = wb.active; ws.title = "صندوق‌ها"; ws.sheet_view.rightToLeft = True
    ws.append([title for _, title in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="18375E")
    for item in items:
        ws.append([item.get(key) for key, _ in columns])
    for col in ws.columns:
        width = min(36, max(12, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[col[0].column_letter].width = width
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=funds.xlsx"})


@app.get("/api/export/funds.json")
def export_funds_json():
    return {"exported_at": utc_now(), "source": "Rahavard365",
            "items": db.funds(source="Rahavard365")}


@app.get("/api/export/analysis.json")
def export_analysis_json(horizon: str = Query(default="long")):
    horizon = normalize_horizon(horizon)
    return {symbol: analyze(symbol, _market_history(symbol), horizon=horizon)
            for symbol in ["GOLD", "DXY", "USD_IR_FREE"] if _market_history(symbol, 1)}
